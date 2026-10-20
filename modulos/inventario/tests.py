import io
from datetime import datetime
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase
from openpyxl import load_workbook

from modulos.inventario.views.informes.saldo_inventario import IndexInformeSaldoInvetarioView


class SaldoInventarioTests(SimpleTestCase):
    def setUp(self):
        self.view = IndexInformeSaldoInvetarioView()
        articulo = SimpleNamespace(codigo='ART-1', descr='Artículo', um='UN', cup=10, tipo='Material')
        self.movimientos = [
            SimpleNamespace(tipo_id=1, tipo=SimpleNamespace(signo=1, nombre='Entrada'),
                            cantidad=10, codigo=articulo, punit=10, fecha=datetime(2026, 1, 1),
                            bodega=1, numero=1),
            SimpleNamespace(tipo_id=12, tipo=SimpleNamespace(signo=1, nombre='Vale de Consumo'),
                            cantidad=3, codigo=articulo, punit=10, fecha=datetime(2026, 1, 2),
                            bodega=1, numero=2),
            SimpleNamespace(tipo_id=2, tipo=SimpleNamespace(signo=-1, nombre='Salida'),
                            cantidad=-2, codigo=articulo, punit=10, fecha=datetime(2026, 1, 3),
                            bodega=1, numero=3),
        ]

    def test_saldo_global_descuenta_vale_de_consumo(self):
        qs = MagicMock()
        qs.iterator.return_value = iter(self.movimientos)
        with patch('modulos.inventario.views.informes.saldo_inventario.Movs.objects') as objects:
            objects.select_related.return_value = qs
            qs.filter.return_value = qs
            qs.exclude.return_value = qs
            qs.distinct.return_value = qs
            qs.order_by.return_value = qs
            datos = self.view._get_datos_globales('', None, True)

        self.assertEqual(len(datos), 1)
        self.assertEqual(datos[0]['saldo'], 5)
        self.assertEqual(datos[0]['costo'], 50)

    def test_excel_articulo_descuenta_vale_de_consumo(self):
        qs = MagicMock()
        qs.iterator.return_value = iter(self.movimientos)
        request = SimpleNamespace(POST={'codigo': 'ART-1', 'fecha_corte': ''})
        with (patch.object(self.view, '_get_movimientos', return_value=(qs, None)),
              patch('modulos.inventario.views.informes.saldo_inventario.Bodegas.objects') as bodegas):
            bodegas.values.return_value = [{'cod': 1, 'nombre': 'Bodega'}]
            response = self.view._generar_excel(request)

        ws = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(ws['B3'].value, -3)
        self.assertEqual(ws['G3'].value, -30)
        self.assertEqual(ws['B3'].font.color.rgb, '00B91C1C')
        self.assertEqual(ws['E3'].fill.fgColor.rgb, '00FEF2F2')
        self.assertIsNone(ws['B2'].font.color)
        self.assertIn('Total Entradas: 10  |  Total Salidas: 5  |  Saldo Final: 5', ws['A6'].value)

    def test_excel_global_destaca_saldo_negativo(self):
        datos = [dict(codigo='ART-1', nombre='Artículo', um='UN', saldo=-3,
                      ultimo_precio=10, costo=-30, tipo_nombre='Material')]
        request = SimpleNamespace(POST={'fecha_corte': '', 'tipo': '', 'solo_con_stock': 'false'})
        with patch.object(self.view, '_get_datos_globales', return_value=datos):
            response = self.view._generar_excel_global(request)

        ws = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(ws['D3'].value, -3)
        self.assertEqual(ws['D3'].font.color.rgb, '00B91C1C')
        self.assertEqual(ws['F3'].fill.fgColor.rgb, '00FEF2F2')

    def test_vale_ya_negativo_no_se_suma_al_saldo(self):
        self.movimientos[1].cantidad = -3
        self.assertEqual(self.view._cantidad_movimiento(self.movimientos[1]), -3)
